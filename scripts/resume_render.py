#!/usr/bin/env python3

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any, Iterable
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RT
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate

from ekb_paths import config_path


class ResumeError(Exception):
    pass


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ResumeError(f"cannot read JSON {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ResumeError(f"{path} must contain a JSON object")
    return value


def whole_item_link_allowed(path: str) -> bool:
    """Whether a sourced item's entire visible text may be a hyperlink.

    Whole-item links are reserved for contact/profile values and entity labels.
    Summary text is always plain. Other narrative prose must name the linked
    entity explicitly through ``link_text`` so a project URL cannot turn an
    entire achievement bullet into blue underlined text.
    """
    patterns = (
        r"basics\.(?:contact|links)\[\d+\]",
        r"experience\[\d+\]\.organization",
        r"experience\[\d+\]\.engagements\[\d+\]\.name",
        r"projects\[\d+\]\.primary",
        r"(?:education|certifications|awards|activities)\[\d+\]\.primary",
    )
    return any(re.fullmatch(pattern, path) for pattern in patterns)


def sourced(value: Any, path: str, nullable: bool = False) -> dict[str, str] | None:
    if value is None and nullable:
        return None
    allowed = {"text", "source_ref", "source_refs", "url", "link_text"}
    if not isinstance(value, dict) or set(value) - allowed:
        raise ResumeError(
            f"{path} must contain text and source_ref or source_refs, and may contain url and link_text"
        )
    has_one = "source_ref" in value
    has_many = "source_refs" in value
    if has_one == has_many:
        raise ResumeError(f"{path} must use exactly one of source_ref or source_refs")
    if "text" not in value or not isinstance(value["text"], str) or not value["text"].strip():
        raise ResumeError(f"{path}.text must be a non-empty string")
    if has_one:
        if not isinstance(value["source_ref"], str) or not value["source_ref"].strip():
            raise ResumeError(f"{path}.source_ref must be a non-empty string")
    else:
        # Composition, for a claim that is inherently plural. The Ruby source
        # checker enforces the evidence rules; this is the structural gate.
        refs = value["source_refs"]
        if not isinstance(refs, list) or len(refs) < 2:
            raise ResumeError(f"{path}.source_refs must list at least two source IDs")
        if len(refs) > 4:
            raise ResumeError(f"{path}.source_refs composes more than four sources")
        if len(set(refs)) != len(refs):
            raise ResumeError(f"{path}.source_refs repeats a source")
        if not all(isinstance(ref, str) and ref.strip() for ref in refs):
            raise ResumeError(f"{path}.source_refs entries must be non-empty strings")
    if "url" in value:
        url = value["url"]
        if not isinstance(url, str) or not url.strip():
            raise ResumeError(f"{path}.url must be a non-empty string when present")
        if not re.match(r"^(https://|mailto:|tel:)", url.strip()):
            raise ResumeError(f"{path}.url must start with https://, mailto:, or tel:")
        if re.fullmatch(r"summary\[\d+\]", path):
            raise ResumeError(f"{path}.url is not allowed; summary prose must render without hyperlinks")
    if "link_text" in value:
        link_text = value["link_text"]
        if "url" not in value:
            raise ResumeError(f"{path}.link_text requires url")
        if not isinstance(link_text, str) or not link_text.strip():
            raise ResumeError(f"{path}.link_text must be a non-empty string when present")
        if value["text"].count(link_text) != 1:
            raise ResumeError(f"{path}.link_text must occur exactly once in text")
        if link_text == value["text"]:
            raise ResumeError(f"{path}.link_text must identify only part of text")
    elif "url" in value and not whole_item_link_allowed(path):
        raise ResumeError(
            f"{path}.url would hyperlink narrative prose; add link_text for the named entity or remove url"
        )
    return value


def url_of(value: dict[str, str] | None) -> str:
    """Optional hyperlink target for a sourced item; empty when it is plain text."""
    return "" if value is None else str(value.get("url", "")).strip()


def link_text_of(value: dict[str, str] | None) -> str:
    """Optional substring that alone receives the item's hyperlink."""
    return "" if value is None else str(value.get("link_text", ""))


def link_fragments(value: dict[str, str]) -> list[tuple[str, bool]]:
    """Split sourced text into plain and linked fragments."""
    text = text_of(value)
    link_text = link_text_of(value)
    if not link_text:
        return [(text, True)]
    before, after = text.split(link_text, 1)
    return [(before, False), (link_text, True), (after, False)]


def links_enabled(model: dict[str, Any]) -> bool:
    return model["layout"].get("hyperlinks", "auto") != "off"


def link_of(model: dict[str, Any], value: dict[str, str] | None) -> str:
    """Hyperlink target honoring the document-level hyperlink mode."""
    return url_of(value) if links_enabled(model) else ""


def link_appearance(policy: dict[str, Any]) -> tuple[str, bool]:
    """Return the explicit cross-format hyperlink treatment."""
    links = policy.get("links", {})
    color = str(links.get("color_hex", "0563C1")).lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", color):
        raise ResumeError("policy.links.color_hex must contain exactly six hexadecimal digits")
    return color, bool(links.get("underline", True))


def bullet_layout(policy: dict[str, Any]) -> dict[str, float]:
    """Validated cross-format bullet geometry in points."""
    raw = policy.get("bullets", {})
    defaults = {
        "text_indent_pt": 20.0,
        "hanging_pt": 11.0,
        "marker_font_size_pt": 6.0,
        "marker_vertical_offset_pt": -1.5,
        # Second level, for named engagements inside a role. The title sits
        # between the employer margin and its own bullets so the reader sees
        # employer, project, achievement as three distinct depths.
        "engagement_title_indent_pt": 20.0,
        "engagement_text_indent_pt": 40.0,
    }
    layout: dict[str, float] = {}
    for key, fallback in defaults.items():
        value = raw.get(key, fallback)
        if not isinstance(value, (int, float)):
            raise ResumeError(f"policy.bullets.{key} must be numeric")
        layout[key] = float(value)
    if layout["text_indent_pt"] <= 0:
        raise ResumeError("policy.bullets.text_indent_pt must be positive")
    if not 0 < layout["hanging_pt"] < layout["text_indent_pt"]:
        raise ResumeError(
            "policy.bullets.hanging_pt must be positive and smaller than text_indent_pt"
        )
    if layout["marker_font_size_pt"] <= 0:
        raise ResumeError("policy.bullets.marker_font_size_pt must be positive")
    if not layout["text_indent_pt"] <= layout["engagement_title_indent_pt"] + layout["hanging_pt"]:
        raise ResumeError(
            "policy.bullets.engagement_title_indent_pt must keep an engagement title "
            "inside the role bullet margin"
        )
    if layout["engagement_text_indent_pt"] <= layout["engagement_title_indent_pt"]:
        raise ResumeError(
            "policy.bullets.engagement_text_indent_pt must be deeper than "
            "engagement_title_indent_pt"
        )
    return layout


def engagements_of(entry: dict[str, Any]) -> list[dict[str, Any]]:
    """Named engagements inside a role, or an empty list for a flat role."""
    value = entry.get("engagements")
    return value if isinstance(value, list) else []


def experience_items(model: dict[str, Any]) -> list[dict[str, str]]:
    """Every sourced achievement in Experience, role bullets and engagement
    bullets alike. Emphasis density, bullet-opening variety, and extraction
    order all need the complete list, not just the top level."""
    items: list[dict[str, str]] = []
    for entry in model.get("experience", []):
        items.extend(entry.get("bullets", []))
        for engagement in engagements_of(entry):
            items.extend(engagement.get("bullets", []))
    return items


def emphasis_plan(model: dict[str, Any], policy: dict[str, Any]) -> tuple[list[str], list[str]]:
    """Terms to bold, derived from the requirements the resume actually matched.

    Deriving them instead of accepting a hand-written list keeps two promises at
    once: a bolded term is always a term the target asked for, and it always has
    selected evidence behind it — the source checker already refuses a `matched`
    requirement without both.

    The policy caps how many requirements may be highlighted, because past
    roughly a dozen nothing stands out any more. The cap selects the most
    important requirements rather than failing the render: emphasis is
    presentation, and no document should be blocked by a formatting preference.
    Returns the selected terms and the requirement terms the cap dropped, which
    are reported rather than silently discarded.
    """
    if model["layout"].get("emphasis", "matched-requirements") == "none":
        return [], []
    candidates = [
        requirement for requirement in model["alignment"]["requirements"]
        if requirement["status"] == "matched" and requirement.get("emphasize", True)
    ]
    # Required-and-critical first, then required, then preferred; stable within
    # each group so the model's own ordering still decides near-equal terms.
    candidates.sort(
        key=lambda requirement: (
            0 if requirement["critical"] and requirement["priority"] == "required"
            else 1 if requirement["priority"] == "required"
            else 2
        )
    )
    maximum = policy.get("emphasis", {}).get("max_terms")
    kept = candidates[:maximum] if maximum else candidates
    dropped = [requirement["term"] for requirement in candidates[len(kept):]]

    terms: list[str] = []
    seen: set[str] = set()
    excluded = {
        str(term).strip().casefold()
        for term in policy.get("emphasis", {}).get("excluded_terms", [])
        if str(term).strip()
    }
    for requirement in kept:
        for term in [requirement["term"], *requirement["aliases"]]:
            cleaned = term.strip()
            folded = cleaned.casefold()
            if cleaned and folded not in excluded and folded not in seen:
                seen.add(folded)
                terms.append(cleaned)
    # Longest first so "Clean Architecture" wins over a nested "architecture".
    return sorted(terms, key=len, reverse=True), dropped


def emphasis_terms(model: dict[str, Any], policy: dict[str, Any]) -> list[str]:
    return emphasis_plan(model, policy)[0]


class Emphasizer:
    """Bolds each term at most once per section, per the ATS policy.

    Highlighting runs after the prose exists, never during drafting: it is a
    reading aid layered onto unchanged words. Bolding every occurrence would
    defeat its own purpose, so the first occurrence in each section wins and the
    rest render normally.
    """

    def __init__(self, terms: Iterable[str], policy: dict[str, Any]) -> None:
        self.terms = list(terms)
        self.limit = int(policy.get("emphasis", {}).get("occurrences_per_term_per_section", 1))
        self.sections = set(policy.get("emphasis", {}).get("sections", []))
        self.used: dict[tuple[str, str], int] = {}
        self.per_section: dict[str, int] = {}
        self.section = ""
        self.count = 0

    def enter(self, section: str) -> None:
        self.section = section

    def active(self) -> bool:
        return bool(self.terms) and self.section in self.sections

    def split(self, text: str) -> list[tuple[str, bool]]:
        """Split text into (fragment, bold) pairs in reading order."""
        if not self.active() or not text:
            return [(text, False)]
        spans: list[tuple[int, int]] = []
        for term in self.terms:
            key = (self.section, term.casefold())
            remaining = self.limit - self.used.get(key, 0)
            if remaining <= 0:
                continue
            for match in re.finditer(rf"(?<![\w.+#-]){re.escape(term)}(?![\w+#-])", text, re.IGNORECASE):
                if any(match.start() < end and start < match.end() for start, end in spans):
                    continue
                spans.append((match.start(), match.end()))
                self.used[key] = self.used.get(key, 0) + 1
                self.per_section[self.section] = self.per_section.get(self.section, 0) + 1
                self.count += 1
                remaining -= 1
                if remaining <= 0:
                    break
        if not spans:
            return [(text, False)]
        parts: list[tuple[str, bool]] = []
        cursor = 0
        for start, end in sorted(spans):
            if start > cursor:
                parts.append((text[cursor:start], False))
            parts.append((text[start:end], True))
            cursor = end
        if cursor < len(text):
            parts.append((text[cursor:], False))
        return parts

    def markup(self, text: str) -> str:
        """The same split expressed as ReportLab inline markup."""
        return "".join(
            f"<b>{html.escape(fragment)}</b>" if bold else html.escape(fragment)
            for fragment, bold in self.split(text)
        )


def sourced_list(value: Any, path: str, minimum: int = 0) -> list[dict[str, str]]:
    if not isinstance(value, list) or len(value) < minimum:
        raise ResumeError(f"{path} must be an array with at least {minimum} item(s)")
    return [sourced(item, f"{path}[{index}]") for index, item in enumerate(value)]  # type: ignore[list-item]


# Default summary length, in words. Four to six sentences need room for the
# supported identity, relevant experience and expertise, career direction, and
# applicable mobility context. The ceiling keeps that overview compact; it is
# not permission to pad beyond the available evidence.
SUMMARY_DEFAULT_WORDS = 90
SUMMARY_MINIMUM_SENTENCES = 4
SUMMARY_MAXIMUM_SENTENCES = 6
SUMMARY_EXPERIENCE_PATTERN = re.compile(
    r"\b(\d+)\+\s+years?\s+of\s+experience\b", flags=re.IGNORECASE
)


def summary_limit(model: dict[str, Any], policy: dict[str, Any] | None = None) -> int:
    """Word limit for the summary: model override, then policy, then default."""
    if "summary_word_limit" in model["layout"]:
        return int(model["layout"]["summary_word_limit"])
    if policy:
        return int(policy.get("summary", {}).get("default_words", SUMMARY_DEFAULT_WORDS))
    return SUMMARY_DEFAULT_WORDS


def summary_sentence_count(text: str) -> int:
    """Count complete prose sentences using terminal punctuation.

    Resume summaries should avoid abbreviations that make sentence boundaries
    ambiguous. Requiring terminal punctuation also catches fragments that would
    otherwise satisfy a raw item count.
    """
    return len(re.findall(r"[.!?]+(?=(?:[\"')\]]*)\s|$)", text.strip()))


def public_resume_text(model: dict[str, Any]) -> list[tuple[str, str]]:
    """Return only text that renders publicly, excluding private targeting data."""
    sections = (
        "basics", "summary", "experience", "projects", "career_breaks", "skills",
        "languages", "education", "certifications", "awards", "activities", "hobbies",
    )
    results: list[tuple[str, str]] = []

    def visit(value: Any, path: str) -> None:
        if isinstance(value, dict):
            if isinstance(value.get("text"), str):
                results.append((f"{path}.text", value["text"]))
                return
            for key, child in value.items():
                if path.startswith("skills[") and key == "name" and isinstance(child, str):
                    results.append((f"{path}.name", child))
                elif isinstance(child, (dict, list)):
                    visit(child, f"{path}.{key}")
        elif isinstance(value, list):
            for index, child in enumerate(value):
                visit(child, f"{path}[{index}]")

    for section in sections:
        if section in model:
            visit(model[section], section)
    return results


def validate_writing_style(model: dict[str, Any], policy: dict[str, Any] | None) -> None:
    if not policy:
        return
    style = policy.get("writing_style", {})
    symbols = style.get("forbidden_symbols", {})
    terms = style.get("forbidden_terms", [])
    phrases = style.get("forbidden_phrases", [])
    if not isinstance(symbols, dict):
        raise ResumeError("policy writing_style.forbidden_symbols must be an object")
    if not isinstance(terms, list) or not all(isinstance(item, str) and item for item in terms):
        raise ResumeError("policy writing_style.forbidden_terms must contain non-empty strings")
    if not isinstance(phrases, list) or not all(isinstance(item, str) and item for item in phrases):
        raise ResumeError("policy writing_style.forbidden_phrases must contain non-empty strings")

    for path, text in public_resume_text(model):
        for symbol, label in symbols.items():
            if symbol in text:
                raise ResumeError(f"{path} contains forbidden {label} symbol {symbol!r}")
        for phrase in [*terms, *phrases]:
            if re.search(rf"(?<!\w){re.escape(phrase)}(?!\w)", text, flags=re.IGNORECASE):
                raise ResumeError(f"{path} contains forbidden AI-style wording {phrase!r}")


def validate_quantifier_quality(model: dict[str, Any], policy: dict[str, Any] | None) -> None:
    """Reject test and source-code size metrics that can be inflated mechanically."""
    if not policy:
        return
    quality = policy.get("quantifier_quality", {})
    context_terms = quality.get("test_context_terms", [])
    units = quality.get("forbidden_test_count_units", [])
    if not isinstance(context_terms, list) or not all(
        isinstance(item, str) and item for item in context_terms
    ):
        raise ResumeError("policy quantifier_quality.test_context_terms must contain non-empty strings")
    if not isinstance(units, list) or not all(isinstance(item, str) and item for item in units):
        raise ResumeError(
            "policy quantifier_quality.forbidden_test_count_units must contain non-empty strings"
        )
    if not context_terms or not units:
        test_enabled = False
    else:
        test_enabled = True

    context_pattern = re.compile(
        r"(?<!\w)(?:" + "|".join(re.escape(item) for item in context_terms) + r")(?!\w)",
        flags=re.IGNORECASE,
    )
    count_pattern = re.compile(
        r"(?<!\w)\d[\d,]*(?:\.\d+)?\s*(?:\+|to|-)?\s*"
        r"(?:" + "|".join(re.escape(item) for item in units) + r")(?!\w)",
        flags=re.IGNORECASE,
    )
    for path, text in public_resume_text(model):
        match = count_pattern.search(text) if test_enabled and context_pattern.search(text) else None
        if match is not None:
            raise ResumeError(
                f"{path} contains a gameable testing-size count {match.group(0)!r}; "
                "describe tested behaviors, boundaries, or failure paths instead"
            )

    code_context_terms = quality.get("code_context_terms", [])
    code_units = quality.get("forbidden_code_count_units", [])
    if not isinstance(code_context_terms, list) or not all(
        isinstance(item, str) and item for item in code_context_terms
    ):
        raise ResumeError("policy quantifier_quality.code_context_terms must contain non-empty strings")
    if not isinstance(code_units, list) or not all(
        isinstance(item, str) and item for item in code_units
    ):
        raise ResumeError(
            "policy quantifier_quality.forbidden_code_count_units must contain non-empty strings"
        )
    if not code_context_terms or not code_units:
        return

    code_context_pattern = re.compile(
        r"(?<!\w)(?:" + "|".join(re.escape(item) for item in code_context_terms) + r")(?!\w)",
        flags=re.IGNORECASE,
    )
    code_count_pattern = re.compile(
        r"(?<!\w)\d[\d,]*(?:\.\d+)?\s*(?:\+)?"
        r"(?:\s*(?:to|of|-)\s*\d[\d,]*(?:\.\d+)?\s*(?:\+)?)?\s*[- ]?\s*"
        r"(?:" + "|".join(re.escape(item) for item in code_units) + r")(?!\w)",
        flags=re.IGNORECASE,
    )
    for path, text in public_resume_text(model):
        if not code_context_pattern.search(text):
            continue
        match = code_count_pattern.search(text)
        if match:
            raise ResumeError(
                f"{path} contains a gameable code-size count {match.group(0)!r}; "
                "describe the structural change or delivered capability instead"
            )


def validate_model(model: dict[str, Any], policy: dict[str, Any] | None = None) -> None:
    required = {
        "schema_version", "application_id", "target", "layout", "basics", "experience",
        "projects", "skills", "education", "certifications", "alignment", "quantifier_review",
    }
    optional = {"summary", "languages", "awards", "activities", "hobbies", "career_breaks"}
    missing, unexpected = required - set(model), set(model) - required - optional
    if missing or unexpected:
        raise ResumeError(f"resume model keys invalid; missing={sorted(missing)}, unexpected={sorted(unexpected)}")
    if model["schema_version"] != 1:
        raise ResumeError("resume model schema_version must be 1")
    app_id = model["application_id"]
    if not isinstance(app_id, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}-[a-z0-9]+(?:-[a-z0-9]+)*", app_id):
        raise ResumeError("application_id is not safe or date-prefixed")
    target = model["target"]
    target_required = {
        "mode", "company", "role", "market", "market_basis", "job_country", "job_country_code",
        "location_scope", "location_basis",
    }
    target_optional = {"summary_lead"}
    if (
        not isinstance(target, dict)
        or not target_required.issubset(target)
        or set(target) - target_required - target_optional
    ):
        raise ResumeError(
            "target must contain mode, company, role, market, market_basis, job_country, job_country_code, "
            "location_scope, and location_basis, "
            "and may contain summary_lead"
        )
    for key in target_required - {"job_country", "job_country_code"}:
        if not isinstance(target[key], str) or not target[key].strip():
            raise ResumeError(f"target.{key} must be a non-empty string")
    if target["job_country"] is not None and (
        not isinstance(target["job_country"], str) or not target["job_country"].strip()
    ):
        raise ResumeError("target.job_country must be null or a non-empty string")
    if target["job_country_code"] is not None and (
        not isinstance(target["job_country_code"], str)
        or not re.fullmatch(r"[A-Z]{2}", target["job_country_code"])
    ):
        raise ResumeError("target.job_country_code must be null or an ISO 3166-1 alpha-2 code")
    if (target["job_country"] is None) != (target["job_country_code"] is None):
        raise ResumeError("target.job_country and target.job_country_code must both be set or both be null")
    if target["mode"] not in {"master", "job-targeted"}:
        raise ResumeError("target.mode must be master or job-targeted")
    if target["market"] not in {"europe", "north-america"}:
        raise ResumeError("target.market must be europe or north-america")
    if target["market_basis"] not in {"job-location", "profile-default", "user-override"}:
        raise ResumeError("target.market_basis is invalid")
    if target["location_scope"] not in {
        "same-country", "outside-country", "location-independent", "unspecified"
    }:
        raise ResumeError("target.location_scope is invalid")
    if target["location_basis"] not in {
        "job-description", "job-url", "user-request", "unresolved"
    }:
        raise ResumeError("target.location_basis is invalid")
    layout = model["layout"]
    layout_keys = {"page_target", "summary_word_limit", "page_break_before", "hyperlinks", "emphasis"}
    if (
        not isinstance(layout, dict)
        or "page_target" not in layout
        or set(layout) - layout_keys
        or layout["page_target"] not in {1, 2}
    ):
        raise ResumeError("layout must contain page_target of 1 or 2 and may contain summary_word_limit")
    if layout.get("hyperlinks", "auto") not in {"auto", "off"}:
        raise ResumeError("layout.hyperlinks must be auto or off")
    if layout.get("emphasis", "matched-requirements") not in {"matched-requirements", "none"}:
        raise ResumeError("layout.emphasis must be matched-requirements or none")
    summary_policy = (policy or {}).get("summary", {})
    minimum_override = int(summary_policy.get("minimum_override", 40))
    maximum_override = int(summary_policy.get("maximum_override", 120))
    if "summary_word_limit" in layout and (
        not isinstance(layout["summary_word_limit"], int)
        or not minimum_override <= layout["summary_word_limit"] <= maximum_override
    ):
        raise ResumeError(
            "layout.summary_word_limit must be an integer from "
            f"{minimum_override} to {maximum_override}"
        )
    if "page_break_before" in layout and layout["page_break_before"] not in {"selected-projects"}:
        raise ResumeError("layout.page_break_before must be selected-projects when present")
    if "page_break_before" in layout and layout["page_target"] != 2:
        raise ResumeError("layout.page_break_before is only valid when page_target is 2")

    basics = model["basics"]
    if not isinstance(basics, dict) or set(basics) != {"name", "contact", "links", "mobility"}:
        raise ResumeError("basics must contain exactly name, contact, links, and mobility")
    sourced(basics["name"], "basics.name")
    sourced_list(basics["contact"], "basics.contact", 1)
    sourced_list(basics["links"], "basics.links")
    sourced(basics["mobility"], "basics.mobility", nullable=True)
    mobility_required = (
        target["mode"] == "master" or target["location_scope"] == "outside-country"
    )
    if mobility_required and basics["mobility"] is None:
        raise ResumeError(
            "basics.mobility is required for master resumes and outside-country jobs"
        )
    if basics["mobility"] is not None:
        mobility_refs = (
            [basics["mobility"].get("source_ref")]
            if basics["mobility"].get("source_ref")
            else list(basics["mobility"].get("source_refs") or [])
        )
        if not any(str(ref).startswith("profile-eligibility-") for ref in mobility_refs):
            raise ResumeError(
                "basics.mobility must cite a profile-eligibility relocation source"
            )
    summary = sourced_list(model.get("summary", []), "summary")
    summary_text = " ".join(text_of(item) for item in summary)
    limit = summary_limit(model, policy)
    if summary and len(re.findall(r"\b[\w’'-]+\b", summary_text)) > limit:
        raise ResumeError(f"summary exceeds the {limit}-word limit")
    if summary:
        minimum_sentences = int(
            summary_policy.get("minimum_sentences", SUMMARY_MINIMUM_SENTENCES)
        )
        maximum_sentences = int(
            summary_policy.get("maximum_sentences", SUMMARY_MAXIMUM_SENTENCES)
        )
        if minimum_sentences < 1 or maximum_sentences < minimum_sentences:
            raise ResumeError(
                "policy.summary sentence bounds must be positive and ordered"
            )
        sentence_count = summary_sentence_count(summary_text)
        if not minimum_sentences <= sentence_count <= maximum_sentences:
            raise ResumeError(
                "summary must contain "
                f"{minimum_sentences} to {maximum_sentences} complete sentences; "
                f"found {sentence_count}"
            )
        claims: list[tuple[dict[str, Any], re.Match[str]]] = []
        for item in summary:
            claims.extend((item, match) for match in SUMMARY_EXPERIENCE_PATTERN.finditer(text_of(item)))
        if len(claims) != 1:
            raise ResumeError(
                "summary must contain exactly one N+ years of experience figure"
            )
        claim_item, claim_match = claims[0]
        first_sentence_end = re.search(r"[.!?]", summary_text)
        if first_sentence_end and claim_match.group(0) not in summary_text[: first_sentence_end.end()]:
            raise ResumeError("the years-of-experience figure must appear in the first summary sentence")
        refs = (
            [claim_item.get("source_ref")]
            if claim_item.get("source_ref")
            else list(claim_item.get("source_refs") or [])
        )
        if not any(str(ref).startswith("profile-experience-") for ref in refs):
            raise ResumeError(
                "the years-of-experience figure must cite at least one profile-experience source"
            )
    if target["mode"] == "job-targeted":
        lead = target.get("summary_lead")
        if not lead:
            raise ResumeError("job-targeted resumes require target.summary_lead")
        if not summary:
            raise ResumeError("job-targeted resumes with target.summary_lead require a summary")
        if not summary_text.casefold().startswith(lead.casefold()):
            raise ResumeError(
                "job-targeted summary must begin with target.summary_lead "
                f"({lead!r})"
            )

    experience = model["experience"]
    if not isinstance(experience, list) or not experience:
        raise ResumeError("experience must contain at least one entry")
    experience_keys = {"organization", "title", "location", "start", "end", "bullets"}
    budget = (policy or {}).get("experience_budget", {})
    engagements_max = int(budget.get("engagements_max", 4))
    engagement_bullets_max = int(budget.get("engagement_bullets_max", 2))
    for index, entry in enumerate(experience):
        path = f"experience[{index}]"
        if not isinstance(entry, dict) or not experience_keys <= set(entry):
            raise ResumeError(f"{path} has invalid keys")
        if set(entry) - experience_keys - {"engagements"}:
            raise ResumeError(f"{path} has invalid keys")
        for key in ("organization", "title", "start", "end"):
            sourced(entry[key], f"{path}.{key}")
        sourced(entry["location"], f"{path}.location", nullable=True)
        bullets = sourced_list(entry["bullets"], f"{path}.bullets", 1)
        if len(bullets) > 4:
            raise ResumeError(f"{path} has more than four bullets")

        # Named engagements. Structural bounds only; how many is an editorial
        # decision surfaced through warnings, because the whole point of this
        # structure is to stop policy from truncating a dense role.
        if "engagements" in entry:
            engagements = entry["engagements"]
            if not isinstance(engagements, list) or not engagements:
                raise ResumeError(f"{path}.engagements must be a non-empty array when present")
            if len(engagements) > engagements_max + 1:
                raise ResumeError(
                    f"{path} has {len(engagements)} engagements; "
                    f"{engagements_max} is the editorial maximum and "
                    f"{engagements_max + 1} the structural limit"
                )
            for position, engagement in enumerate(engagements):
                sub = f"{path}.engagements[{position}]"
                if not isinstance(engagement, dict) or not {"name", "bullets"} <= set(engagement):
                    raise ResumeError(f"{sub} must contain name and bullets")
                if set(engagement) - {"name", "context", "bullets"}:
                    raise ResumeError(f"{sub} has invalid keys")
                sourced(engagement["name"], f"{sub}.name")
                # A name is a label, not a sentence. Overlong ones only surfaced
                # as an unexplained page overflow, so say what actually happened.
                if len(text_of(engagement["name"])) > 60:
                    raise ResumeError(
                        f"{sub}.name is {len(text_of(engagement['name']))} characters; an "
                        "engagement name is a short project label, not a bullet"
                    )
                sourced(engagement.get("context"), f"{sub}.context", nullable=True)
                sub_bullets = sourced_list(engagement["bullets"], f"{sub}.bullets", 1)
                if len(sub_bullets) > engagement_bullets_max + 1:
                    raise ResumeError(
                        f"{sub} has {len(sub_bullets)} bullets; "
                        f"{engagement_bullets_max} is the editorial maximum"
                    )

    if not isinstance(model["skills"], list):
        raise ResumeError("skills must be an array")
    for index, group in enumerate(model["skills"]):
        if not isinstance(group, dict) or set(group) != {"name", "items"} or not isinstance(group["name"], str) or not group["name"].strip():
            raise ResumeError(f"skills[{index}] must contain a non-empty name and items")
        sourced_list(group["items"], f"skills[{index}].items", 1)

    entry_keys = {"primary", "secondary", "date", "details"}
    minimum_projects = int(
        (policy or {}).get("selected_projects", {}).get("minimum_distinct_projects", 2)
    )
    if not isinstance(model["projects"], list) or len(model["projects"]) < minimum_projects:
        raise ResumeError(
            f"projects must contain at least {minimum_projects} Selected Projects entries"
        )
    for section in ("projects", "education", "certifications", "awards", "activities"):
        entries = model.get(section, [])
        if not isinstance(entries, list):
            raise ResumeError(f"{section} must be an array")
        for index, entry in enumerate(entries):
            path = f"{section}[{index}]"
            if not isinstance(entry, dict) or set(entry) != entry_keys:
                raise ResumeError(f"{path} has invalid keys")
            sourced(entry["primary"], f"{path}.primary")
            sourced(entry["secondary"], f"{path}.secondary", nullable=True)
            sourced(entry["date"], f"{path}.date", nullable=True)
            sourced_list(entry["details"], f"{path}.details")
    for section in ("languages", "hobbies"):
        sourced_list(model.get(section, []), section)
    for index, entry in enumerate(model.get("career_breaks", [])):
        path = f"career_breaks[{index}]"
        if not isinstance(entry, dict) or set(entry) != {"label", "start", "end", "details"}:
            raise ResumeError(f"{path} has invalid keys")
        for key in ("label", "start", "end"):
            sourced(entry[key], f"{path}.{key}")
        sourced_list(entry["details"], f"{path}.details")

    alignment = model["alignment"]
    if not isinstance(alignment, dict) or set(alignment) != {"requirements"} or not isinstance(alignment["requirements"], list):
        raise ResumeError("alignment must contain requirements")
    for index, requirement in enumerate(alignment["requirements"]):
        path = f"alignment.requirements[{index}]"
        required_keys = {"term", "aliases", "priority", "critical", "status", "evidence_refs"}
        optional_keys = {"emphasize", "selection_reason"}
        if not isinstance(requirement, dict) or set(requirement) - optional_keys != required_keys:
            raise ResumeError(f"{path} has invalid keys")
        if "emphasize" in requirement and not isinstance(requirement["emphasize"], bool):
            raise ResumeError(f"{path}.emphasize must be true or false")
        if "selection_reason" in requirement and (
            not isinstance(requirement["selection_reason"], str)
            or not requirement["selection_reason"].strip()
        ):
            raise ResumeError(f"{path}.selection_reason must be a non-empty string")
        if not isinstance(requirement["term"], str) or not requirement["term"].strip():
            raise ResumeError(f"{path}.term must be non-empty")
        if not isinstance(requirement["aliases"], list) or not all(isinstance(item, str) and item.strip() for item in requirement["aliases"]):
            raise ResumeError(f"{path}.aliases must be non-empty strings")
        if requirement["priority"] not in {"required", "preferred"} or not isinstance(requirement["critical"], bool):
            raise ResumeError(f"{path} priority or critical value is invalid")
        if requirement["status"] not in {"matched", "supported_not_selected", "unsupported"}:
            raise ResumeError(f"{path}.status is invalid")
        if not isinstance(requirement["evidence_refs"], list) or not all(isinstance(item, str) and item.strip() for item in requirement["evidence_refs"]):
            raise ResumeError(f"{path}.evidence_refs must be an array of IDs")
        if requirement["status"] == "unsupported" and requirement["evidence_refs"]:
            raise ResumeError(f"{path} cannot give evidence_refs for an unsupported requirement")
        if requirement["status"] != "unsupported" and not requirement["evidence_refs"]:
            raise ResumeError(f"{path} needs evidence_refs when it is supported")
        # `selection_reason` records why a selection decision was made, and that
        # applies to every status. An earlier version banned it on `unsupported`
        # on the theory that there is nothing to explain when no evidence exists.
        # That was wrong, and it deadlocked against the understatement gate: a
        # requirement can be `unsupported` BECAUSE the available evidence was
        # judged insufficient to claim it, which is precisely the decision most
        # worth recording. The Capgemini banking-domain requirement is the case
        # that exposed it. The checker demanded a reason; this rule forbade one.
        #
        # No status restriction remains. The field is validated as a non-empty
        # string above, and a reason on a genuinely empty requirement is
        # harmless documentation rather than a defect.
        if (
            target["mode"] == "job-targeted"
            and requirement["priority"] == "required"
            and requirement["status"] == "supported_not_selected"
            and "selection_reason" not in requirement
        ):
            raise ResumeError(
                f"{path} is a supported required job requirement and needs selection_reason"
            )
    quantifiers = model["quantifier_review"]
    if not isinstance(quantifiers, dict) or set(quantifiers) != {"used", "available_not_used", "opportunities"}:
        raise ResumeError("quantifier_review must contain used, available_not_used, and opportunities")
    for key in quantifiers:
        if not isinstance(quantifiers[key], list) or not all(isinstance(item, str) and item.strip() for item in quantifiers[key]):
            raise ResumeError(f"quantifier_review.{key} must be an array of non-empty strings")
    validate_writing_style(model, policy)
    validate_quantifier_quality(model, policy)


def run_source_check(
    model: Path,
    profile: Path,
    projects: Path,
    ranking: Path | None = None,
    index: Path | None = None,
) -> dict[str, Any]:
    checker = Path(__file__).with_name("resume_source_check.rb")
    command = ["ruby", str(checker), "--model", str(model), "--profile", str(profile), "--projects", str(projects)]
    if ranking is not None and ranking.is_file():
        command += ["--ranking", str(ranking)]
    # The evidence index powers the understatement gate. Default to the one
    # beside the repository root so callers do not have to pass it.
    if index is None:
        index = profile.parent.parent / "index" / "evidence-index.yaml"
    if index.is_file():
        command += ["--index", str(index)]
    rubric = Path(config_path("review-rubric.json"))
    if rubric.is_file():
        command += ["--rubric", str(rubric)]
    process = subprocess.run(
        command,
        text=True,
        capture_output=True,
        check=False,
    )
    if not process.stdout.strip():
        raise ResumeError(f"source validation failed: {process.stderr.strip() or 'no report produced'}")
    try:
        report = json.loads(process.stdout)
    except json.JSONDecodeError as exc:
        raise ResumeError(f"source validator returned invalid JSON: {exc}") from exc
    if process.returncode not in (0, 1):
        raise ResumeError(process.stderr.strip() or "source validator could not run")
    return report


def coverage_section(source_report: dict[str, Any]) -> str:
    """One-line coverage summary for the validation report header.

    `demonstrated` means an achievement shows the capability; `stated` means the
    resume only claims it, usually through the Skills list. The distinction is
    what a reviewer scores, so it belongs above the fold."""
    counts = source_report.get("coverage", {}).get("counts") or {}
    if not counts:
        return ""
    order = ["demonstrated", "stated", "unsupported"]
    parts = [f"{counts[key]} {key}" for key in order if counts.get(key)]
    line = f"- Requirement coverage: {', '.join(parts)}\n"
    if not source_report.get("checks", {}).get("evidence_index_loaded", True):
        line += "- Evidence index: not loaded (understatement gate inactive)\n"
    return line


def selection_score_section(source_report: dict[str, Any]) -> str:
    """Private selection score against `config/review-rubric.json`.

    Scores how well the resume used the evidence available to it, NOT the
    candidate's fit. Requirements with no eligible evidence are excluded, so an
    honest gap cannot lower it, and it is not a prediction of any external
    reviewer's number."""
    score = source_report.get("selection_score")
    if not score:
        return ""
    parts = []
    for name, entry in (score.get("dimensions") or {}).items():
        value = entry.get("value")
        parts.append(f"{name} {value if value is not None else 'n/a'}")
    line = (
        f"- Selection score: {score['score']}/100 "
        f"(rubric v{score.get('rubric_version')}; "
        f"{score.get('scoreable_required')} scoreable required requirements, "
        f"{score.get('excluded_no_evidence')} excluded as genuine gaps)\n"
    )
    if parts:
        line += f"  - {', '.join(parts)}\n"
    return line


def external_signals_section(source_report: dict[str, Any]) -> str:
    """What the separate resume-review prompt is likely to penalize.

    Reported apart from the selection score on purpose. The selection score
    excludes requirements the knowledge base cannot answer, because those are
    not the toolkit's failure; the external reviewer weighs them anyway, so its
    blind spot is exactly what belongs here."""
    signals = source_report.get("external_signals")
    if not signals:
        return ""
    lines = ["\n## Likely external review findings\n"]
    seniority = signals.get("seniority")
    if seniority:
        verdict = "meets" if seniority["status"] == "met" else "below"
        lines.append(
            f"- Seniority: posting asks {seniority['required_years']}+ years, "
            f"confirmed timeline is {seniority['confirmed_years']} years ({verdict}). "
            "State this plainly; never narrow it in the document.\n"
        )
    if signals.get("unsupported_critical"):
        lines.append(
            f"- {signals['unsupported_critical']} critical requirement(s) have no eligible evidence.\n"
        )
    gap = signals.get("keyword_gap") or []
    if gap:
        fixable = signals.get("keyword_gap_fixable", 0)
        lines.append(
            f"\nKeyword gap ({len(gap)} terms, {fixable} fixable from existing evidence). "
            "A reviewer's missing-keyword list will draw from these:\n\n"
        )
        lines.append("| Term | Coverage | Priority | Fixable here |\n| --- | --- | --- | --- |\n")
        for entry in gap:
            priority = entry["priority"] + (", critical" if entry.get("critical") else "")
            fix = "yes, evidence exists" if entry.get("fixable") else "no, real gap"
            lines.append(f"| {entry['term']} | {entry['coverage']} | {priority} | {fix} |\n")
    return "".join(lines)


def coverage_table(source_report: dict[str, Any]) -> str:
    """Per-requirement detail.

    Critical stated capabilities show their strongest unused project evidence.
    Non-critical stated tools remain visible as `claimed` without pressuring the
    draft to turn a routine matching record into an achievement.
    """
    requirements = source_report.get("coverage", {}).get("requirements") or []
    if not requirements:
        return ""
    rows = []
    for entry in requirements:
        unused = ", ".join(entry.get("unused_candidates") or []) or "-"
        reason = entry.get("selection_reason") or ""
        marker = {"demonstrated": "shown", "stated": "claimed", "unsupported": "absent"}.get(
            entry.get("coverage", ""), entry.get("coverage", "")
        )
        rows.append(
            f"| {entry.get('term', '')} | {entry.get('priority', '')} | {marker} | "
            f"{unused} | {reason} |"
        )
    return (
        "\n## Requirement coverage\n\n"
        "| Requirement | Priority | Coverage | Strongest unused | Reason |\n"
        "| --- | --- | --- | --- | --- |\n" + "\n".join(rows) + "\n"
    )


def text_of(value: dict[str, str] | None) -> str:
    return "" if value is None else value["text"].strip()


def resume_lines(model: dict[str, Any]) -> list[str]:
    basics = model["basics"]
    lines = [text_of(basics["name"])]
    contact = [text_of(item) for item in basics["contact"] + basics["links"]]
    if contact:
        lines.append(" | ".join(contact))
    if basics["mobility"] is not None:
        lines.append(text_of(basics["mobility"]))
    summary = model.get("summary", [])
    if summary:
        lines.append("SUMMARY")
        lines.append(" ".join(text_of(item) for item in summary))
    lines.append("EXPERIENCE")
    for entry in model["experience"]:
        lines.append(" | ".join(part for part in (text_of(entry["organization"]), text_of(entry["title"])) if part))
        date_location = " | ".join(
            part
            for part in (f"{text_of(entry['start'])} - {text_of(entry['end'])}", text_of(entry["location"]))
            if part
        )
        lines.append(date_location)
        lines.extend(f"• {text_of(item)}" for item in entry["bullets"])
        for engagement in engagements_of(entry):
            lines.append(
                " | ".join(
                    part
                    for part in (text_of(engagement["name"]), text_of(engagement.get("context")))
                    if part
                )
            )
            lines.extend(f"• {text_of(item)}" for item in engagement["bullets"])
    if model.get("projects"):
        lines.append("SELECTED PROJECTS")
        for entry in model["projects"]:
            main = " | ".join(
                part
                for part in (text_of(entry["primary"]), text_of(entry["secondary"]), text_of(entry["date"]))
                if part
            )
            lines.append(main)
            lines.extend(f"• {text_of(item)}" for item in entry["details"])
    if model.get("career_breaks"):
        lines.append("CAREER BREAKS")
        for entry in model["career_breaks"]:
            lines.append(text_of(entry["label"]))
            lines.append(f"{text_of(entry['start'])} - {text_of(entry['end'])}")
            lines.extend(text_of(item) for item in entry["details"])
    if model["skills"]:
        lines.append("SKILLS")
        for group in model["skills"]:
            lines.append(f"{group['name']}: {', '.join(text_of(item) for item in group['items'])}")
    if model.get("languages"):
        lines.append("LANGUAGES")
        lines.append(", ".join(text_of(item) for item in model["languages"]))
    for key, heading in (("education", "EDUCATION"), ("certifications", "CERTIFICATIONS"), ("awards", "AWARDS"), ("activities", "ACTIVITIES")):
        if model.get(key):
            lines.append(heading)
            for entry in model[key]:
                main = " | ".join(part for part in (text_of(entry["primary"]), text_of(entry["secondary"]), text_of(entry["date"])) if part)
                lines.append(main)
                lines.extend(text_of(item) for item in entry["details"])
    if model.get("hobbies"):
        lines.append("INTERESTS")
        lines.append(", ".join(text_of(item) for item in model["hobbies"]))
    return [line for line in lines if line]


def text_export_lines(model: dict[str, Any], policy: dict[str, Any]) -> list[str]:
    """Plain-text export. TXT has no annotation layer, so a hyperlink either
    becomes a visible address or is lost; `policy.links.text_export` decides.
    """
    if policy.get("links", {}).get("text_export") != "append-url" or not links_enabled(model):
        return resume_lines(model)

    def bare(value: str) -> str:
        return re.sub(r"^(https://|mailto:|tel:)", "", normalize(value)).rstrip("/")

    targets: list[tuple[str, str]] = []
    def collect(value: Any) -> None:
        if isinstance(value, dict):
            if "text" in value and ("source_ref" in value or "source_refs" in value):
                url = url_of(value)
                label = link_text_of(value) or text_of(value)
                # Skip anything whose visible text already is the address. This
                # remains useful for custom policies even though the shipped
                # policy keeps profile addresses behind labels.
                if url and bare(url) != bare(label):
                    targets.append((label, url))
                return
            for child in value.values():
                collect(child)
        elif isinstance(value, list):
            for child in value:
                collect(child)

    collect(model)
    lines = resume_lines(model)
    for label, url in targets:
        needle = normalize(label)
        for index, line in enumerate(lines):
            if needle and needle in normalize(line) and url not in line:
                lines[index] = re.sub(
                    rf"({re.escape(label)})(?! \()", rf"\1 ({url})", line, count=1
                )
                break
    return lines


def page_size(model: dict[str, Any], policy: dict[str, Any]) -> tuple[float, float]:
    sizes = policy.get("page_sizes", {})
    market = model["target"]["market"]
    name = sizes.get(market)
    if name == "A4":
        return A4
    if name == "LETTER":
        return LETTER
    raise ResumeError(f"policy has no page size for {market}")


def continuation_contact(model: dict[str, Any]) -> str:
    basics = model["basics"]
    values = [text_of(basics["name"])] + [text_of(item) for item in basics["contact"] + basics["links"]]
    return " | ".join(value for value in values if value)


def add_docx_hyperlink(
    paragraph: Any,
    text: str,
    url: str,
    font: str,
    size: float,
    bold: bool = False,
    color: str = "0563C1",
    underline: bool = True,
) -> None:
    """Append a clickable run. python-docx has no hyperlink API, so build the
    OOXML w:hyperlink element and its external relationship directly.

    The visible text is whatever the model supplies; keep it meaningful, because
    text-only ATS parsers read the run text and discard the relationship target.
    """
    r_id = paragraph.part.relate_to(url, DOCX_RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    link_color = OxmlElement("w:color")
    link_color.set(qn("w:val"), color)
    properties.append(link_color)
    if underline:
        link_underline = OxmlElement("w:u")
        link_underline.set(qn("w:val"), "single")
        properties.append(link_underline)
    hyperlink.append(run)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    node.set(qn("xml:space"), "preserve")
    run.append(node)
    paragraph._p.append(hyperlink)
    # Re-apply every visible property rather than depending on a viewer's
    # Hyperlink style, which is not consistent across Word-compatible apps.
    rpr_font = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rpr_font.set(qn(attr), font)
    properties.append(rpr_font)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    properties.append(sz)
    if bold:
        properties.append(OxmlElement("w:b"))


def add_docx_sourced_runs(
    paragraph: Any,
    model: dict[str, Any],
    item: dict[str, str],
    policy: dict[str, Any],
    emphasizer: Emphasizer | None = None,
    bold: bool = False,
) -> None:
    """Render one sourced item, retaining emphasis inside an optional link."""
    font = policy["fonts"]["primary"]
    size = policy["sizes_pt"]["body"]
    target = link_of(model, item)
    color, underline = link_appearance(policy)
    for link_fragment, is_linked in link_fragments(item):
        fragments = emphasizer.split(link_fragment) if emphasizer else [(link_fragment, False)]
        for fragment, emphasized in fragments:
            if not fragment:
                continue
            fragment_bold = bold or emphasized
            if target and is_linked:
                add_docx_hyperlink(
                    paragraph,
                    fragment,
                    target,
                    font,
                    size,
                    bold=fragment_bold,
                    color=color,
                    underline=underline,
                )
            else:
                set_run_font(paragraph.add_run(fragment), font, size, bold=fragment_bold)


def set_run_font(run: Any, name: str, size: float, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_docx_section_heading(document: Document, label: str, policy: dict[str, Any]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(policy["spacing_pt"]["section_before"])
    paragraph.paragraph_format.space_after = Pt(policy["spacing_pt"]["section_after"])
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(label.upper())
    set_run_font(run, policy["fonts"]["primary"], policy["sizes_pt"]["section"], bold=True)
    paragraph._p.get_or_add_pPr().append(_bottom_border())


def _bottom_border() -> Any:
    from docx.oxml import OxmlElement

    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    borders.append(bottom)
    return borders


def configure_docx_bullet_marker(document: Document, policy: dict[str, Any]) -> None:
    """Lower the real List Bullet numbering glyph to the body-text baseline."""
    layout = bullet_layout(policy)
    numbering = document.part.numbering_part.element
    for abstract_numbering in numbering.findall(qn("w:abstractNum")):
        for level in abstract_numbering.findall(qn("w:lvl")):
            paragraph_style = level.find(qn("w:pStyle"))
            if (
                paragraph_style is None
                or paragraph_style.get(qn("w:val")) != "ListBullet"
            ):
                continue
            run_properties = level.find(qn("w:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("w:rPr")
                level.append(run_properties)
            position = run_properties.find(qn("w:position"))
            if position is None:
                position = OxmlElement("w:position")
                run_properties.append(position)
            position.set(
                qn("w:val"),
                str(round(layout["marker_vertical_offset_pt"] * 2)),
            )
            return
    raise ResumeError("DOCX template does not contain the List Bullet numbering definition")


def render_docx(model: dict[str, Any], policy: dict[str, Any], output: Path) -> Emphasizer:
    emphasizer = Emphasizer(emphasis_terms(model, policy), policy)
    hyperlink_color, hyperlink_underline = link_appearance(policy)
    document = Document()
    section = document.sections[0]
    width, height = page_size(model, policy)
    section.page_width = Inches(width / inch)
    section.page_height = Inches(height / inch)
    margins = policy["margins_inches"]
    section.top_margin = Inches(margins["top"])
    section.right_margin = Inches(margins["right"])
    section.bottom_margin = Inches(margins["bottom"])
    section.left_margin = Inches(margins["left"])
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run(continuation_contact(model)), policy["fonts"]["primary"], policy["sizes_pt"]["small"])

    normal = document.styles["Normal"]
    normal.font.name = policy["fonts"]["primary"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), policy["fonts"]["primary"])
    normal.font.size = Pt(policy["sizes_pt"]["body"])
    normal.paragraph_format.space_after = Pt(policy["spacing_pt"]["paragraph_after"])
    normal.paragraph_format.line_spacing = 1.0
    configure_docx_bullet_marker(document, policy)
    bullets = bullet_layout(policy)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(1)
    set_run_font(name.add_run(text_of(model["basics"]["name"])), policy["fonts"]["primary"], policy["sizes_pt"]["name"], bold=True)

    contact_items = model["basics"]["contact"] + model["basics"]["links"]
    if contact_items:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        font, size = policy["fonts"]["primary"], policy["sizes_pt"]["small"]
        for index, item in enumerate(contact_items):
            if index:
                set_run_font(paragraph.add_run(" | "), font, size)
            target = link_of(model, item)
            if target:
                add_docx_hyperlink(
                    paragraph,
                    text_of(item),
                    target,
                    font,
                    size,
                    color=hyperlink_color,
                    underline=hyperlink_underline,
                )
            else:
                set_run_font(paragraph.add_run(text_of(item)), font, size)

    mobility = model["basics"]["mobility"]
    if mobility is not None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        set_run_font(
            paragraph.add_run(text_of(mobility)),
            policy["fonts"]["primary"],
            policy["sizes_pt"]["small"],
        )

    summary = model.get("summary", [])
    if summary:
        emphasizer.enter("Summary")
        add_docx_section_heading(document, "Summary", policy)
        paragraph = document.add_paragraph()
        for index, item in enumerate(summary):
            if index:
                set_run_font(
                    paragraph.add_run(" "),
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["body"],
                )
            add_docx_sourced_runs(paragraph, model, item, policy, emphasizer)
        paragraph.paragraph_format.space_after = Pt(policy["spacing_pt"]["paragraph_after"])

    emphasizer.enter("Experience")
    add_docx_section_heading(document, "Experience", policy)
    for entry in model["experience"]:
        heading = document.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(0)
        organization_url = link_of(model, entry["organization"])
        if organization_url:
            add_docx_hyperlink(
                heading,
                text_of(entry["organization"]),
                organization_url,
                policy["fonts"]["primary"],
                policy["sizes_pt"]["body"],
                bold=True,
                color=hyperlink_color,
                underline=hyperlink_underline,
            )
        else:
            set_run_font(heading.add_run(text_of(entry["organization"])), policy["fonts"]["primary"], policy["sizes_pt"]["body"], bold=True)
        set_run_font(heading.add_run(f" | {text_of(entry['title'])}"), policy["fonts"]["primary"], policy["sizes_pt"]["body"], bold=True)
        meta = document.add_paragraph()
        meta.paragraph_format.keep_with_next = True
        meta.paragraph_format.space_after = Pt(1)
        date_location = " | ".join(
            part
            for part in (f"{text_of(entry['start'])} - {text_of(entry['end'])}", text_of(entry["location"]))
            if part
        )
        set_run_font(meta.add_run(date_location), policy["fonts"]["primary"], policy["sizes_pt"]["small"])
        for item in entry["bullets"]:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Pt(bullets["text_indent_pt"])
            paragraph.paragraph_format.first_line_indent = Pt(-bullets["hanging_pt"])
            paragraph.paragraph_format.space_after = Pt(policy["spacing_pt"]["bullet_after"])
            add_docx_sourced_runs(paragraph, model, item, policy, emphasizer)

        for engagement in engagements_of(entry):
            title = document.add_paragraph()
            title.paragraph_format.keep_with_next = True
            title.paragraph_format.left_indent = Pt(bullets["engagement_title_indent_pt"])
            title.paragraph_format.space_before = Pt(1)
            title.paragraph_format.space_after = Pt(0)
            name_url = link_of(model, engagement["name"])
            if name_url:
                add_docx_hyperlink(
                    title,
                    text_of(engagement["name"]),
                    name_url,
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["body"],
                    bold=True,
                    color=hyperlink_color,
                    underline=hyperlink_underline,
                )
            else:
                set_run_font(
                    title.add_run(text_of(engagement["name"])),
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["body"],
                    bold=True,
                )
            context = text_of(engagement.get("context"))
            if context:
                set_run_font(
                    title.add_run(f" | {context}"),
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["small"],
                )
            for item in engagement["bullets"]:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Pt(bullets["engagement_text_indent_pt"])
                paragraph.paragraph_format.first_line_indent = Pt(-bullets["hanging_pt"])
                paragraph.paragraph_format.space_after = Pt(policy["spacing_pt"]["bullet_after"])
                add_docx_sourced_runs(paragraph, model, item, policy, emphasizer)

    if model.get("projects"):
        emphasizer.enter("Selected Projects")
        if model["layout"].get("page_break_before") == "selected-projects":
            document.add_page_break()
        add_docx_section_heading(document, "Selected Projects", policy)
        for entry in model["projects"]:
            heading = document.add_paragraph()
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.space_after = Pt(1)
            primary = entry["primary"]
            primary_url = link_of(model, primary)
            if primary_url:
                add_docx_hyperlink(
                    heading,
                    text_of(primary),
                    primary_url,
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["body"],
                    bold=True,
                    color=hyperlink_color,
                    underline=hyperlink_underline,
                )
            else:
                set_run_font(
                    heading.add_run(text_of(primary)),
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["body"],
                    bold=True,
                )
            tail = " | ".join(
                part for part in (text_of(entry["secondary"]), text_of(entry["date"])) if part
            )
            if tail:
                heading.add_run(f" | {tail}")
            for detail in entry["details"]:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Pt(bullets["text_indent_pt"])
                paragraph.paragraph_format.first_line_indent = Pt(-bullets["hanging_pt"])
                paragraph.paragraph_format.space_after = Pt(policy["spacing_pt"]["bullet_after"])
                add_docx_sourced_runs(paragraph, model, detail, policy, emphasizer)

    if model.get("career_breaks"):
        add_docx_section_heading(document, "Career Breaks", policy)
        for entry in model["career_breaks"]:
            heading = document.add_paragraph()
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.space_after = Pt(0)
            set_run_font(heading.add_run(text_of(entry["label"])), policy["fonts"]["primary"], policy["sizes_pt"]["body"], bold=True)
            meta = document.add_paragraph(f"{text_of(entry['start'])} - {text_of(entry['end'])}")
            meta.paragraph_format.keep_with_next = True
            meta.paragraph_format.space_after = Pt(1)
            for detail in entry["details"]:
                document.add_paragraph(text_of(detail))

    if model["skills"]:
        emphasizer.enter("Skills")
        add_docx_section_heading(document, "Skills", policy)
        for group in model["skills"]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(policy["spacing_pt"]["paragraph_after"])
            set_run_font(paragraph.add_run(f"{group['name']}: "), policy["fonts"]["primary"], policy["sizes_pt"]["body"], bold=True)
            # The group label is already bold, so spend each term's one
            # occurrence on it rather than repeating the same word two words
            # later ("Flutter and Dart: Flutter, Dart").
            emphasizer.split(group["name"])
            for index, item in enumerate(group["items"]):
                if index:
                    set_run_font(paragraph.add_run(", "), policy["fonts"]["primary"], policy["sizes_pt"]["body"])
                add_docx_sourced_runs(paragraph, model, item, policy, emphasizer)

    if model.get("languages"):
        add_docx_section_heading(document, "Languages", policy)
        document.add_paragraph(", ".join(text_of(item) for item in model["languages"]))

    for key, heading_label in (("education", "Education"), ("certifications", "Certifications"), ("awards", "Awards"), ("activities", "Activities")):
        if not model.get(key):
            continue
        add_docx_section_heading(document, heading_label, policy)
        for entry in model[key]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = bool(entry["details"])
            primary_url = link_of(model, entry["primary"])
            if primary_url:
                add_docx_hyperlink(
                    paragraph,
                    text_of(entry["primary"]),
                    primary_url,
                    policy["fonts"]["primary"],
                    policy["sizes_pt"]["body"],
                    bold=True,
                    color=hyperlink_color,
                    underline=hyperlink_underline,
                )
            else:
                set_run_font(paragraph.add_run(text_of(entry["primary"])), policy["fonts"]["primary"], policy["sizes_pt"]["body"], bold=True)
            tail = " | ".join(part for part in (text_of(entry["secondary"]), text_of(entry["date"])) if part)
            if tail:
                paragraph.add_run(f" | {tail}")
            for detail in entry["details"]:
                detail_paragraph = document.add_paragraph()
                add_docx_sourced_runs(detail_paragraph, model, detail, policy)

    if model.get("hobbies"):
        add_docx_section_heading(document, "Interests", policy)
        document.add_paragraph(", ".join(text_of(item) for item in model["hobbies"]))

    document.core_properties.title = f"Resume: {text_of(model['basics']['name'])}"
    document.core_properties.subject = f"Application for {model['target']['role']} at {model['target']['company']}"
    document.save(output)
    return emphasizer


def render_pdf(model: dict[str, Any], policy: dict[str, Any], output: Path) -> None:
    emphasizer = Emphasizer(emphasis_terms(model, policy), policy)
    hyperlink_color, hyperlink_underline = link_appearance(policy)
    margins = policy["margins_inches"]
    document = SimpleDocTemplate(
        str(output),
        pagesize=page_size(model, policy),
        rightMargin=margins["right"] * inch,
        leftMargin=margins["left"] * inch,
        topMargin=margins["top"] * inch,
        bottomMargin=margins["bottom"] * inch,
        title=f"Resume: {text_of(model['basics']['name'])}",
        subject=f"Application for {model['target']['role']} at {model['target']['company']}",
    )
    sample = getSampleStyleSheet()
    font = policy["fonts"]["pdf_primary"]
    body_size = policy["sizes_pt"]["body"]
    bullet_tokens = bullet_layout(policy)
    styles = {
        "name": ParagraphStyle("EKBName", parent=sample["Normal"], fontName=f"{font}-Bold", fontSize=policy["sizes_pt"]["name"], leading=policy["sizes_pt"]["name"] + 2, alignment=TA_CENTER, spaceAfter=1),
        "contact": ParagraphStyle("EKBContact", parent=sample["Normal"], fontName=font, fontSize=policy["sizes_pt"]["small"], leading=policy["sizes_pt"]["small"] + 1.5, alignment=TA_CENTER, spaceAfter=2),
        "mobility": ParagraphStyle("EKBMobility", parent=sample["Normal"], fontName=font, fontSize=policy["sizes_pt"]["small"], leading=policy["sizes_pt"]["small"] + 1.5, alignment=TA_CENTER, spaceAfter=5),
        "section": ParagraphStyle("EKBSection", parent=sample["Normal"], fontName=f"{font}-Bold", fontSize=policy["sizes_pt"]["section"], leading=policy["sizes_pt"]["section"] + 1, spaceBefore=policy["spacing_pt"]["section_before"], spaceAfter=policy["spacing_pt"]["section_after"], borderWidth=0, borderPadding=0, keepWithNext=True),
        "entry": ParagraphStyle("EKBEntry", parent=sample["Normal"], fontName=font, fontSize=body_size, leading=body_size + 1.4, spaceAfter=0, keepWithNext=True),
        "body": ParagraphStyle("EKBBody", parent=sample["Normal"], fontName=font, fontSize=body_size, leading=body_size + 1.4, spaceAfter=policy["spacing_pt"]["paragraph_after"]),
        "meta": ParagraphStyle("EKBMeta", parent=sample["Normal"], fontName=font, fontSize=policy["sizes_pt"]["small"], leading=policy["sizes_pt"]["small"] + 1.1, spaceAfter=1, keepWithNext=True),
    }

    def paragraph(text: str, style: str = "body") -> Paragraph:
        return Paragraph(html.escape(text), styles[style])

    def section(label: str) -> list[Any]:
        return [paragraph(label.upper(), "section")]

    story: list[Any] = [paragraph(text_of(model["basics"]["name"]), "name")]

    def anchor(body: str, target: str) -> str:
        if hyperlink_underline:
            body = f"<u>{body}</u>"
        return (
            f'<a href="{html.escape(target, quote=True)}" '
            f'color="#{hyperlink_color}">{body}</a>'
        )

    def linked(item: dict[str, str]) -> str:
        target = link_of(model, item)
        if not target:
            return html.escape(text_of(item))
        return "".join(
            anchor(html.escape(fragment), target) if is_linked else html.escape(fragment)
            for fragment, is_linked in link_fragments(item)
        )

    def linked_emphasized_markup(item: dict[str, str]) -> str:
        target = link_of(model, item)
        return "".join(
            anchor(emphasizer.markup(fragment), target)
            if target and is_linked
            else emphasizer.markup(fragment)
            for fragment, is_linked in link_fragments(item)
        )

    def linked_emphasized(item: dict[str, str], style: str = "body") -> Paragraph:
        return Paragraph(linked_emphasized_markup(item), styles[style])

    contact_items = model["basics"]["contact"] + model["basics"]["links"]
    if contact_items:
        story.append(Paragraph(" | ".join(linked(item) for item in contact_items), styles["contact"]))
    mobility = model["basics"]["mobility"]
    if mobility is not None:
        story.append(paragraph(text_of(mobility), "mobility"))

    summary = model.get("summary", [])
    if summary:
        emphasizer.enter("Summary")
        story.extend(section("Summary"))
        story.append(
            Paragraph(
                " ".join(linked_emphasized_markup(item) for item in summary),
                styles["body"],
            )
        )

    emphasizer.enter("Experience")
    story.extend(section("Experience"))
    for entry in model["experience"]:
        organization = linked(entry["organization"])
        heading = Paragraph(
            f"<b>{organization} | {html.escape(text_of(entry['title']))}</b>",
            styles["entry"],
        )
        date_location = " | ".join(
            part
            for part in (f"{text_of(entry['start'])} - {text_of(entry['end'])}", text_of(entry["location"]))
            if part
        )
        bullets = ListFlowable(
            [ListItem(linked_emphasized(item)) for item in entry["bullets"]],
            bulletType="bullet",
            start="circle",
            leftIndent=bullet_tokens["text_indent_pt"],
            bulletDedent=bullet_tokens["hanging_pt"],
            bulletFontName=font,
            bulletFontSize=bullet_tokens["marker_font_size_pt"],
            bulletOffsetY=bullet_tokens["marker_vertical_offset_pt"],
            spaceAfter=2,
        )
        story.extend([heading, paragraph(date_location, "meta"), bullets])

        for engagement in engagements_of(entry):
            name = linked(engagement["name"])
            context = text_of(engagement.get("context"))
            label = f"<b>{name}</b>"
            if context:
                label += f" | {html.escape(context)}"
            engagement_heading = Paragraph(label, styles["entry"])
            engagement_heading.style = styles["entry"].clone("engagementEntry")
            engagement_heading.style.leftIndent = bullet_tokens["engagement_title_indent_pt"]
            engagement_heading.style.spaceBefore = 1
            story.append(engagement_heading)
            story.append(
                ListFlowable(
                    [ListItem(linked_emphasized(item)) for item in engagement["bullets"]],
                    bulletType="bullet",
                    start="circle",
                    leftIndent=bullet_tokens["engagement_text_indent_pt"],
                    bulletDedent=bullet_tokens["hanging_pt"],
                    bulletFontName=font,
                    bulletFontSize=bullet_tokens["marker_font_size_pt"],
                    bulletOffsetY=bullet_tokens["marker_vertical_offset_pt"],
                    spaceAfter=2,
                )
            )

    if model.get("projects"):
        emphasizer.enter("Selected Projects")
        if model["layout"].get("page_break_before") == "selected-projects":
            story.append(PageBreak())
        story.extend(section("Selected Projects"))
        for entry in model["projects"]:
            primary = linked(entry["primary"])
            tail = " | ".join(
                part for part in (text_of(entry["secondary"]), text_of(entry["date"])) if part
            )
            main = f"<b>{primary}</b>"
            if tail:
                main += f" | {html.escape(tail)}"
            flowables: list[Any] = [Paragraph(main, styles["entry"])]
            if entry["details"]:
                flowables.append(
                    ListFlowable(
                        [ListItem(linked_emphasized(item)) for item in entry["details"]],
                        bulletType="bullet",
                        start="circle",
                        leftIndent=bullet_tokens["text_indent_pt"],
                        bulletDedent=bullet_tokens["hanging_pt"],
                        bulletFontName=font,
                        bulletFontSize=bullet_tokens["marker_font_size_pt"],
                        bulletOffsetY=bullet_tokens["marker_vertical_offset_pt"],
                        spaceAfter=2,
                    )
                )
            story.append(KeepTogether(flowables))

    if model.get("career_breaks"):
        story.extend(section("Career Breaks"))
        for entry in model["career_breaks"]:
            story.extend([
                Paragraph(f"<b>{html.escape(text_of(entry['label']))}</b>", styles["entry"]),
                paragraph(f"{text_of(entry['start'])} - {text_of(entry['end'])}", "meta"),
            ])
            story.extend(paragraph(text_of(item)) for item in entry["details"])

    if model["skills"]:
        emphasizer.enter("Skills")
        story.extend(section("Skills"))
        for group in model["skills"]:
            emphasizer.split(group["name"])  # see the DOCX renderer for why
            items = ", ".join(
                linked_emphasized_markup(item) for item in group["items"]
            )
            story.append(Paragraph(f"<b>{html.escape(group['name'])}:</b> {items}", styles["body"]))

    if model.get("languages"):
        story.extend(section("Languages"))
        story.append(paragraph(", ".join(text_of(item) for item in model["languages"])))

    for key, heading_label in (("education", "Education"), ("certifications", "Certifications"), ("awards", "Awards"), ("activities", "Activities")):
        if not model.get(key):
            continue
        story.extend(section(heading_label))
        for entry in model[key]:
            tail = " | ".join(part for part in (text_of(entry["secondary"]), text_of(entry["date"])) if part)
            main = f"<b>{linked(entry['primary'])}</b>"
            if tail:
                main += f" | {html.escape(tail)}"
            flowables: list[Any] = [Paragraph(main, styles["body"])]
            flowables.extend(Paragraph(linked(item), styles["body"]) for item in entry["details"])
            story.append(KeepTogether(flowables))

    if model.get("hobbies"):
        story.extend(section("Interests"))
        story.append(paragraph(", ".join(text_of(item) for item in model["hobbies"])))

    def later_page(canvas: Any, _document: Any) -> None:
        canvas.saveState()
        header_style = ParagraphStyle(
            "EKBContinuation",
            parent=sample["Normal"],
            fontName=font,
            fontSize=policy["sizes_pt"]["small"],
            leading=policy["sizes_pt"]["small"] + 1,
            alignment=TA_CENTER,
        )
        header = Paragraph(html.escape(continuation_contact(model)), header_style)
        available_width = page_size(model, policy)[0] - (margins["left"] + margins["right"]) * inch
        _, header_height = header.wrap(available_width, 0.45 * inch)
        header.drawOn(
            canvas,
            margins["left"] * inch,
            page_size(model, policy)[1] - 0.18 * inch - header_height,
        )
        canvas.restoreState()

    document.build(story, onFirstPage=lambda _canvas, _document: None, onLaterPages=later_page)


def extract_docx(path: Path) -> tuple[str, dict[str, Any]]:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    header_footer_text: list[str] = []
    for section in document.sections:
        header_footer_text.extend(p.text for p in section.header.paragraphs if p.text.strip())
        header_footer_text.extend(p.text for p in section.footer.paragraphs if p.text.strip())
    structure = {
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "header_footer_text": header_footer_text,
        "sections": len(document.sections),
        "section_types": [str(section.start_type) for section in document.sections],
    }
    return "\n".join(paragraphs), structure


def pdf_content_fill_ratio(page: Any, policy: dict[str, Any]) -> float:
    """Measure the vertical span of visible text inside the configured margins."""
    height = float(page.mediabox.height)
    margins = policy["margins_inches"]
    usable_bottom = float(margins["bottom"]) * 72
    usable_top = height - float(margins["top"]) * 72
    extents: list[tuple[float, float]] = []

    def visitor(text: str, cm: list[float], _tm: list[float], _font: Any, font_size: float) -> None:
        if not text or not text.strip():
            return
        baseline = float(cm[5])
        size = float(font_size or 0)
        if baseline < usable_bottom - 12 or baseline > usable_top + 24:
            return
        extents.append((baseline, baseline + size))

    page.extract_text(visitor_text=visitor)
    if not extents or usable_top <= usable_bottom:
        return 0.0
    content_bottom = max(min(low for low, _high in extents), usable_bottom)
    content_top = min(max(high for _low, high in extents), usable_top)
    return round(max(0.0, min(1.0, (content_top - content_bottom) / (usable_top - usable_bottom))), 3)


def extract_pdf(path: Path, policy: dict[str, Any]) -> tuple[str, int, str, list[float]]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ResumeError("PDF is password-protected; application resumes must open without a password")
    pages = [(page.extract_text() or "") for page in reader.pages]
    first = reader.pages[0].mediabox
    width, height = round(float(first.width)), round(float(first.height))
    detected = "A4" if abs(width - 595) <= 2 and abs(height - 842) <= 2 else "LETTER" if (width, height) == (612, 792) else f"{width}x{height}pt"
    fill_ratios = [pdf_content_fill_ratio(page, policy) for page in reader.pages]
    return "\n".join(pages), len(reader.pages), detected, fill_ratios


def normalize(value: str) -> str:
    value = value.replace("•", " ").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", value).strip().casefold()


def verify_order(expected: Iterable[str], extracted: str, label: str) -> list[str]:
    haystack = normalize(extracted)
    cursor = 0
    errors: list[str] = []
    for fragment in expected:
        needle = normalize(fragment.lstrip("• "))
        position = haystack.find(needle, cursor)
        if position < 0:
            errors.append(f"{label} is missing or reorders: {fragment}")
        else:
            cursor = position + len(needle)
    return errors


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def filename_tokens(value: str) -> list[str]:
    """Portable filename words without guessing or adding resume facts."""
    ascii_value = (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
    )
    return re.findall(r"[A-Za-z0-9]+", ascii_value)


def delivery_stem(model: dict[str, Any], policy: dict[str, Any]) -> str:
    """Human-facing filename recommended for the application attachment.

    The stable ``resume.*`` files remain available for tooling. The additional
    attachment copies make the output directory directly usable without a
    manual rename such as ``CV_FINAL_v3.pdf``.
    """
    name = filename_tokens(text_of(model["basics"]["name"]))
    role = filename_tokens(model["target"]["role"])
    values = {
        "FirstName": name[0] if name else "Resume",
        "LastName": "_".join(name[1:]),
        "TargetRole": "_".join(role) if role else "Role",
    }
    pattern = str(
        policy.get("delivery", {}).get(
            "attachment_pattern", "FirstName_LastName_TargetRole_CV.pdf"
        )
    )
    pattern = re.sub(r"\.(?:pdf|docx)$", "", pattern, flags=re.IGNORECASE)
    for key, value in values.items():
        pattern = pattern.replace(key, value)
    stem = "_".join(filename_tokens(pattern))
    return stem or "Resume_CV"


def count_links(model: dict[str, Any]) -> int:
    """Sourced items that render as a live link in this document."""
    if not links_enabled(model):
        return 0
    total = 0

    def walk(value: Any) -> None:
        nonlocal total
        if isinstance(value, dict):
            if "text" in value and ("source_ref" in value or "source_refs" in value):
                total += 1 if url_of(value) else 0
                return
            for child in value.values():
                walk(child)
        elif isinstance(value, list):
            for child in value:
                walk(child)

    walk(model)
    return total


def link_warnings(model: dict[str, Any]) -> list[str]:
    """Hyperlinking is the default, so a rendered entity with no link is worth
    reporting — the address may simply be missing from the profile."""
    warnings: list[str] = []
    if not links_enabled(model):
        return ["Hyperlink annotations are disabled for this document (layout.hyperlinks=off)"]
    for entry in model["experience"]:
        if not url_of(entry["organization"]):
            warnings.append(f"No hyperlink on employer: {text_of(entry['organization'])}")
        for engagement in engagements_of(entry):
            if not url_of(engagement["name"]):
                warnings.append(f"No hyperlink on engagement: {text_of(engagement['name'])}")
    for entry in model.get("projects", []):
        if not url_of(entry["primary"]):
            warnings.append(f"No hyperlink on selected project: {text_of(entry['primary'])}")
    for entry in model.get("certifications", []):
        if not url_of(entry["primary"]):
            warnings.append(f"No hyperlink on certification: {text_of(entry['primary'])}")
    return warnings


def emphasis_warnings(model: dict[str, Any], policy: dict[str, Any], emphasizer: Emphasizer) -> list[str]:
    warnings: list[str] = []
    matched = [
        requirement for requirement in model["alignment"]["requirements"]
        if requirement["status"] == "matched"
    ]
    if model["layout"].get("emphasis", "matched-requirements") == "none":
        if matched:
            warnings.append(f"Keyword emphasis is off while {len(matched)} requirements are matched")
        return warnings
    if matched and not emphasizer.terms:
        warnings.append("Every matched requirement has emphasize=false, so nothing is highlighted")
    if emphasizer.terms and emphasizer.count == 0:
        warnings.append("No emphasis term was found in the emphasized sections; check the aliases")
    dropped = emphasis_plan(model, policy)[1]
    if dropped:
        warnings.append(
            f"Emphasis capped at {policy['emphasis']['max_terms']} requirements; not highlighted: "
            + ", ".join(dropped)
        )
    # Skills is a list of terms, so bolding there is inherently dense and says
    # nothing about readability. Judge density on the prose sections only.
    prose = sum(
        count for section, count in emphasizer.per_section.items()
        if section in {"Summary", "Experience", "Selected Projects"}
    )
    bullets = len(experience_items(model)) + sum(
        len(entry["details"]) for entry in model.get("projects", [])
    )
    if bullets and prose > 1.5 * bullets:
        warnings.append(
            f"{prose} emphasized fragments across {bullets} bullets reads as heavily bolded; "
            "reduce the term set"
        )
    return warnings


def application_guidance_warnings(model: dict[str, Any], policy: dict[str, Any]) -> list[str]:
    """Report presentation guidance that cannot safely be auto-corrected.

    Contact types and language levels are user facts. A renderer may point out
    a missing convention, but it must never invent a phone number, classify an
    employer email, or translate a proficiency label into CEFR by itself.
    """
    warnings: list[str] = []
    contact_policy = policy.get("contact_hygiene", {})
    contact = model["basics"]["contact"]
    links = model["basics"]["links"]
    visible_contact = [text_of(item).strip() for item in contact]
    urls = [url_of(item) for item in contact]

    if contact_policy.get("require_email", True) and not any(
        "@" in text or url.startswith("mailto:")
        for text, url in zip(visible_contact, urls)
    ):
        warnings.append("Header has no recognizable professional email address")
    if contact_policy.get("recommend_phone", True) and not any(
        url.startswith("tel:") or len(re.sub(r"\D", "", text)) >= 7
        for text, url in zip(visible_contact, urls)
    ):
        warnings.append("Header has no phone number; add one only from confirmed profile data")
    if contact_policy.get("recommend_profile_link", True) and not any(
        url_of(item) for item in links
    ):
        warnings.append(
            "Header has no confirmed professional profile link; never guess or construct one"
        )

    language_policy = policy.get("languages", {})
    cefr_pattern = language_policy.get("cefr_pattern", r"\b(?:A1|A2|B1|B2|C1|C2)\b")
    if language_policy.get("prefer_cefr", True):
        try:
            cefr = re.compile(str(cefr_pattern), flags=re.IGNORECASE)
        except re.error as exc:
            raise ResumeError(f"policy.languages.cefr_pattern is invalid: {exc}") from exc
        for index, item in enumerate(model.get("languages", []), 1):
            if not cefr.search(text_of(item)):
                warnings.append(
                    f"Language entry {index} has no CEFR level (A1-C2); preserve the recorded "
                    "proficiency unless the user confirms a CEFR mapping"
                )
    return warnings


def editorial_warnings(model: dict[str, Any]) -> list[str]:
    visible = " ".join(resume_lines(model)).casefold()
    warnings: list[str] = []
    for phrase in ("excellent problem solver", "results-driven professional", "proven track record", "dynamic professional", "passionate engineer"):
        if phrase in visible:
            warnings.append(f"Avoid generic resume phrase: {phrase}")
    all_bullets = [text_of(item) for item in experience_items(model)]
    for index, entry in enumerate(model["experience"], 1):
        count = len(entry["bullets"])
        current = text_of(entry["end"]).casefold() == "present"
        engagements = engagements_of(entry)
        if engagements:
            # The role delegates its detail to named engagements, so the
            # role-level bullets are framing lines and the flat targets do not
            # apply. What matters here is that the framing stays short and the
            # engagements stay readable.
            if count > 2:
                warnings.append(
                    f"Role {index} has {count} role-level bullets alongside "
                    f"{len(engagements)} engagements; one or two framing lines is the target"
                )
            if len(engagements) > 4:
                warnings.append(
                    f"Role {index} has {len(engagements)} engagements; four is the default maximum"
                )
            for position, engagement in enumerate(engagements, 1):
                if len(engagement["bullets"]) > 2:
                    warnings.append(
                        f"Role {index} engagement {position} "
                        f"({text_of(engagement['name'])}) has {len(engagement['bullets'])} "
                        "bullets; two is the default maximum"
                    )
        elif current and count != 3:
            warnings.append(f"Current role {index} has {count} bullets; three is the default target")
        elif not current and count > 2:
            warnings.append(f"Historical role {index} has {count} bullets; one or two is the default target")
    starts = [re.sub(r"^[^a-z]*([a-z]+).*", r"\1", bullet.casefold()) for bullet in all_bullets]
    repeated = sorted({verb for verb in starts if verb and starts.count(verb) > 2})
    if repeated:
        warnings.append("Repeated bullet openings: " + ", ".join(repeated))
    for index, bullet in enumerate(all_bullets, 1):
        if bullet and bullet[-1] not in ".!?":
            warnings.append(f"Bullet {index} does not end with consistent punctuation")
    return warnings


def content_density_warnings(
    model: dict[str, Any],
    policy: dict[str, Any],
    pdf_pages: int,
    fill_ratios: list[float],
) -> list[str]:
    if pdf_pages != 1 or not fill_ratios:
        return []
    threshold = float(
        policy.get("content_density", {}).get("single_page_minimum_usable_height_ratio", 0)
    )
    ratio = fill_ratios[0]
    if threshold and ratio < threshold:
        return [
            f"Single-page resume is underfilled: content uses {ratio:.1%} of printable height, "
            f"below the {threshold:.1%} minimum; add target-relevant Selected Projects "
            "evidence or another distinct project"
        ]
    return []


def allowed_continuation_header(model: dict[str, Any], texts: list[str]) -> bool:
    expected = normalize(continuation_contact(model))
    return all(normalize(text) == expected for text in texts)


def apply_plain_mode(model: dict[str, Any]) -> None:
    """Strip both presentation layers for a maximally conservative export.

    Neither hyperlinks nor bold changes a single visible word, so a plain export
    is the same resume with the annotations removed — never a different document.
    """
    model["layout"]["hyperlinks"] = "off"
    model["layout"]["emphasis"] = "none"


def render(args: argparse.Namespace) -> int:
    model_path = Path(args.model).resolve()
    profile_path = Path(args.profile).resolve()
    projects_path = Path(args.projects).resolve()
    policy_path = Path(args.policy).resolve() if args.policy else Path(config_path("resume-policy.json"))
    output_dir = Path(args.output_dir).resolve()
    ranking_path = Path(args.ranking).resolve() if args.ranking else profile_path.with_name("project-ranking.yaml")

    model = load_json(model_path)
    policy = load_json(policy_path)
    validate_model(model, policy)
    if args.ats_plain:
        apply_plain_mode(model)
    expected_policy = model["schema_version"]
    if policy.get("version") != expected_policy:
        raise ResumeError(f"policy version must be {expected_policy} for this resume model")
    if len(set(policy.get("sizes_pt", {}).values())) > 2:
        raise ResumeError("ATS policy must use no more than two font sizes")
    source_report = run_source_check(model_path, profile_path, projects_path, ranking_path)
    if not source_report.get("valid"):
        raise ResumeError("source validation failed:\n- " + "\n- ".join(source_report.get("errors", [])))

    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ekb-resume-", dir=output_dir.parent) as temporary:
        temp_dir = Path(temporary)
        docx_path = temp_dir / "resume.docx"
        pdf_path = temp_dir / "resume.pdf"
        emphasizer = render_docx(model, policy, docx_path)
        render_pdf(model, policy, pdf_path)

        expected = resume_lines(model)
        docx_text, docx_structure = extract_docx(docx_path)
        pdf_text, pdf_pages, detected_page_size, fill_ratios = extract_pdf(pdf_path, policy)
        errors: list[str] = []
        warnings = editorial_warnings(model)
        warnings.extend(application_guidance_warnings(model, policy))
        warnings.extend(content_density_warnings(model, policy, pdf_pages, fill_ratios))
        warnings.extend(link_warnings(model))
        warnings.extend(emphasis_warnings(model, policy, emphasizer))
        warnings.extend(source_report.get("warnings", []))
        errors.extend(verify_order(expected, docx_text, "DOCX"))
        errors.extend(verify_order(expected, pdf_text, "PDF"))
        if docx_structure["tables"]:
            errors.append("DOCX contains tables")
        if docx_structure["inline_shapes"]:
            errors.append("DOCX contains inline images or shapes")
        if not allowed_continuation_header(model, docx_structure["header_footer_text"]):
            errors.append("DOCX contains non-permitted factual header/footer text")
        if len(normalize(pdf_text)) < 50:
            errors.append("PDF does not contain enough selectable text")
        expected_page_size = "A4" if model["target"]["market"] == "europe" else "LETTER"
        if detected_page_size != expected_page_size:
            errors.append(f"PDF page size is {detected_page_size}, expected {expected_page_size}")
        if pdf_pages > 2:
            errors.append("PDF exceeds the two-page maximum")
        if pdf_pages > model["layout"]["page_target"]:
            warnings.append(f"PDF exceeds the preferred {model['layout']['page_target']}-page target")

        attachment_stem = delivery_stem(model, policy)
        attachment_docx = temp_dir / f"{attachment_stem}.docx"
        attachment_pdf = temp_dir / f"{attachment_stem}.pdf"
        shutil.copyfile(docx_path, attachment_docx)
        shutil.copyfile(pdf_path, attachment_pdf)
        outputs = {
            "resume.docx": docx_path,
            "resume.pdf": pdf_path,
            attachment_docx.name: attachment_docx,
            attachment_pdf.name: attachment_pdf,
        }
        if args.include_text:
            text_path = temp_dir / "resume.txt"
            text_path.write_text("\n".join(text_export_lines(model, policy)) + "\n", encoding="utf-8")
            outputs["resume.txt"] = text_path

        report = {
            "valid": not errors,
            "policy_version": policy["version"],
            "application_id": model["application_id"],
            "resume_mode": model["target"]["mode"],
            "source_validation": source_report,
            "presentation": {
                "hyperlinks": model["layout"].get("hyperlinks", "auto"),
                "hyperlink_color": f"#{link_appearance(policy)[0]}",
                "hyperlink_underline": link_appearance(policy)[1],
                "emphasis": model["layout"].get("emphasis", "matched-requirements"),
                "emphasis_terms": emphasizer.terms,
                "emphasis_applied": emphasizer.count,
                "linked_items": count_links(model),
                "bullet_layout": bullet_layout(policy),
                "attachment_stem": attachment_stem,
            },
            "document_validation": {
                "errors": errors,
                "warnings": warnings,
                "docx_structure": docx_structure,
                "pdf_pages": pdf_pages,
                "page_size": detected_page_size,
                "pdf_password_protected": False,
                "content_fill_ratios": fill_ratios,
                "expected_lines": len(expected),
            },
            "outputs": {
                name: {"bytes": path.stat().st_size, "sha256": sha256(path)}
                for name, path in outputs.items()
            },
        }
        validation_json = temp_dir / "validation.json"
        validation_json.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        validation_md = temp_dir / "validation.md"
        validation_md.write_text(
            "# Resume validation\n\n"
            f"- Status: {'PASS' if report['valid'] else 'FAIL'}\n"
            f"- Application: `{model['application_id']}`\n"
            f"- Resume mode: {report['resume_mode']}\n"
            f"- ATS policy: v{policy['version']}\n"
            f"- Recommended attachment: `{attachment_pdf.name}`\n"
            f"- Selected sources: {', '.join(source_report['selected_sources'])}\n"
            f"- Hyperlinks: {report['presentation']['hyperlinks']} "
            f"({report['presentation']['linked_items']} linked items; "
            f"{report['presentation']['hyperlink_color']}, "
            f"{'underlined' if report['presentation']['hyperlink_underline'] else 'not underlined'})\n"
            f"- Emphasis: {report['presentation']['emphasis']} "
            f"({len(emphasizer.terms)} terms, {emphasizer.count} applied)\n"
            + (
                f"- Project evidence used anywhere: "
                f"{', '.join(source_report['project_selection']['evidence_used'])}"
                f" (ranks {', '.join(str(rank) for rank in source_report['project_selection']['evidence_ranks'])})\n"
                if source_report.get("project_selection", {}).get("evidence_used")
                else ""
            )
            + (
                f"- Named in Selected Projects: "
                f"{', '.join(source_report['project_selection']['named_in_selected_projects'])}"
                f" (ranks {', '.join(str(rank) for rank in source_report['project_selection']['named_ranks'])})\n"
                if source_report.get("project_selection", {}).get("named_in_selected_projects")
                else "- Named in Selected Projects: none\n"
            )
            + coverage_section(source_report)
            + selection_score_section(source_report)
            + f"- PDF pages: {pdf_pages}\n"
            f"- Content fill by page: {', '.join(f'{ratio:.1%}' for ratio in fill_ratios)}\n"
            f"- DOCX tables: {docx_structure['tables']}\n"
            f"- DOCX images/shapes: {docx_structure['inline_shapes']}\n"
            f"- Extraction/order errors: {len(errors)}\n"
            f"- Editorial/layout warnings: {len(warnings)}\n"
            + coverage_table(source_report)
            + external_signals_section(source_report)
            + (("\n## Warnings\n\n" + "\n".join(f"- {warning}" for warning in warnings) + "\n") if warnings else "")
            + (("\n## Errors\n\n" + "\n".join(f"- {error}" for error in errors) + "\n") if errors else ""),
            encoding="utf-8",
        )
        outputs["validation.json"] = validation_json
        outputs["validation.md"] = validation_md

        if errors:
            raise ResumeError("document validation failed:\n- " + "\n- ".join(errors))

        for name, path in outputs.items():
            os.replace(path, output_dir / name)

    print(json.dumps({"valid": True, "output_dir": str(output_dir), "selected_sources": source_report["selected_sources"]}, indent=2))
    return 0


def validate(args: argparse.Namespace) -> int:
    model_path = Path(args.model).resolve()
    profile_path = Path(args.profile).resolve()
    model = load_json(model_path)
    policy_path = (
        Path(args.policy).resolve()
        if args.policy
        else Path(config_path("resume-policy.json"))
    )
    policy = load_json(policy_path)
    validate_model(model, policy)
    ranking_path = Path(args.ranking).resolve() if args.ranking else profile_path.with_name("project-ranking.yaml")
    report = run_source_check(model_path, profile_path, Path(args.projects).resolve(), ranking_path)
    print(json.dumps(report, indent=2))
    return 0 if report.get("valid") else 1


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description="Validate and render evidence-linked EKB resumes")
    commands = root.add_subparsers(dest="command", required=True)
    for name in ("validate", "render"):
        command = commands.add_parser(name)
        command.add_argument("--model", required=True)
        command.add_argument("--profile", required=True)
        command.add_argument("--projects", required=True)
        command.add_argument(
            "--ranking",
            help="global project ranking; defaults to project-ranking.yaml beside the profile",
        )
        if name == "render":
            command.add_argument(
                "--policy",
                help="writing and structure policy; defaults to config/resume-policy.json or its workspace override",
            )
            command.add_argument("--output-dir", required=True)
            command.add_argument("--include-text", action="store_true")
            command.add_argument(
                "--ats-plain",
                action="store_true",
                help="render the same text with no hyperlink annotations and no bold emphasis",
            )
        else:
            command.add_argument(
                "--policy",
                help="writing and structure policy; defaults to config/resume-policy.json",
            )
    return root


def main() -> int:
    args = parser().parse_args()
    try:
        return render(args) if args.command == "render" else validate(args)
    except ResumeError as exc:
        print(f"EKB Resume: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
